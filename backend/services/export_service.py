import io
import re
from bs4 import BeautifulSoup
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import htmldocx


def setup_academic_document(topic: str = "") -> docx.Document:
    """
    Khởi tạo Document Word với thể thức chuẩn văn bản học thuật Việt Nam:
    - Khổ giấy A4 (210mm x 297mm)
    - Lề chuẩn: Trên 2cm, Dưới 2cm, Trái 3cm, Phải 2cm
    - Font Times New Roman 13pt, giãn dòng 1.5 lines
    """
    doc = docx.Document()

    # Khổ giấy A4 và lề trang chuẩn
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    # Cấu hình Style Normal
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Times New Roman"
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)
    style_normal.paragraph_format.space_before = Pt(0)

    # Đảm bảo các Heading cũng dùng Times New Roman
    for heading_name, pt_size in [("Heading 1", 15), ("Heading 2", 14), ("Heading 3", 13)]:
        try:
            h_style = doc.styles[heading_name]
            h_style.font.name = "Times New Roman"
            h_style.font.size = Pt(pt_size)
            h_style.font.bold = True
            h_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            h_style.paragraph_format.space_before = Pt(12)
            h_style.paragraph_format.space_after = Pt(6)
            h_style.paragraph_format.line_spacing = 1.3
        except KeyError:
            pass

    return doc


def export_to_docx(topic: str, html_content: str) -> io.BytesIO:
    """
    Chuyển đổi nội dung HTML từ Tiptap Editor thành file Word .docx
    giữ nguyên các định dạng học thuật (H1, H2, H3, bold, italic, lists, bibliography, tables).
    """
    doc = setup_academic_document(topic)

    # Thêm Tiêu đề Đề tài ở trang đầu nếu có
    if topic:
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_before = Pt(12)
        title_p.paragraph_format.space_after = Pt(20)
        run = title_p.add_run(topic.upper())
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    # Chuẩn bị HTML sạch để nạp vào htmldocx
    soup = BeautifulSoup(html_content or "", "html.parser")

    # Xử lý các thẻ citation-entry để thụt lề treo (hanging indent) nếu cần
    for p in soup.find_all("p", class_="citation-entry"):
        p["style"] = "margin-left: 20px; text-indent: -20px;"

    clean_html = str(soup)

    # Sử dụng HtmlToDocx parser
    parser = htmldocx.HtmlToDocx()
    parser.add_html_to_document(clean_html, doc)

    # Ghi vào BytesIO buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_to_markdown(topic: str, html_content: str) -> str:
    """
    Chuyển đổi nội dung HTML sang định dạng Markdown chuẩn GitHub Flavored Markdown (GFM).
    """
    soup = BeautifulSoup(html_content or "", "html.parser")

    lines = []
    if topic:
        lines.append(f"# {topic.strip()}\n")

    def process_node(node) -> str:
        if isinstance(node, str):
            return node

        name = node.name.lower() if node.name else ""

        # Headings
        if name == "h1":
            return f"\n# {node.get_text().strip()}\n"
        elif name == "h2":
            return f"\n## {node.get_text().strip()}\n"
        elif name == "h3":
            return f"\n### {node.get_text().strip()}\n"
        elif name == "h4":
            return f"\n#### {node.get_text().strip()}\n"

        # Paragraphs & Citation
        elif name == "p":
            inner = "".join(process_node(child) for child in node.children)
            if "citation-entry" in node.get("class", []):
                return f"\n> {inner.strip()}\n"
            return f"\n{inner.strip()}\n"

        # Formats
        elif name in ("strong", "b"):
            inner = "".join(process_node(child) for child in node.children)
            return f"**{inner}**"
        elif name in ("em", "i"):
            inner = "".join(process_node(child) for child in node.children)
            return f"*{inner}*"
        elif name == "u":
            inner = "".join(process_node(child) for child in node.children)
            return f"<u>{inner}</u>"
        elif name == "a":
            href = node.get("href", "")
            text = node.get_text() or href
            return f"[{text}]({href})"
        elif name == "blockquote":
            inner = "".join(process_node(child) for child in node.children)
            return f"\n> {inner.strip()}\n"

        # Lists
        elif name == "ul":
            items = []
            for li in node.find_all("li", recursive=False):
                li_text = "".join(process_node(child) for child in li.children).strip()
                items.append(f"- {li_text}")
            return "\n" + "\n".join(items) + "\n"
        elif name == "ol":
            items = []
            for idx, li in enumerate(node.find_all("li", recursive=False)):
                li_text = "".join(process_node(child) for child in li.children).strip()
                items.append(f"{idx + 1}. {li_text}")
            return "\n" + "\n".join(items) + "\n"

        # Fallback for containers
        else:
            return "".join(process_node(child) for child in node.children)

    body_text = process_node(soup)
    lines.append(body_text)

    full_md = "\n".join(lines)
    # Rút gọn các dòng trống liên tiếp
    full_md = re.sub(r"\n{3,}", "\n\n", full_md).strip() + "\n"
    return full_md
